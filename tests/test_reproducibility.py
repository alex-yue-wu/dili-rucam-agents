import json
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook
import pytest

from dili_rucam_agents.batch import PdfRunFailure, PdfRunResult, _main
from dili_rucam_agents.reproducibility import (
    ReproducibilityRunRecord,
    _score_statistics,
    _write_reproducibility_workbook,
    run_reproducibility_folder,
    validate_repeats,
)


def _complete_report(total_score: int = 6) -> str:
    payload = {
        "injury_pattern": "hepatocellular",
        "R_ratio": 6.4,
        "rucam_scores": {
            "time_to_onset": 2,
            "course": 1,
            "risk_factors": 0,
            "concomitant_drugs": 0,
            "other_causes_excluded": 2,
            "known_hepatotoxicity": 1,
            "rechallenge": 0,
        },
        "total_score": total_score,
        "category": "Probable",
    }
    return (
        "## SECTION A\n\nClinical summary\n\n"
        f"## SECTION B\n\n| Item | Score |\n| --- | --- |\n| Total | {total_score} |\n\n"
        f"## SECTION C\n\n```json\n{json.dumps(payload)}\n```\n"
    )


def _write_complete_reports(result_dir: Path) -> None:
    result_dir.mkdir(parents=True, exist_ok=True)
    for report_name in (
        "analyst-alpha_report.md",
        "analyst-beta_report.md",
        "analyst-gamma_report.md",
    ):
        (result_dir / report_name).write_text(_complete_report(), encoding="utf-8")


def test_score_statistics_computes_sample_spread_and_agreement():
    result = _score_statistics([6, 6, 8, 8, 8])

    assert result.valid_repeats == 5
    assert result.mean == pytest.approx(7.2)
    assert result.sample_standard_deviation == pytest.approx(1.095445115)
    assert result.minimum == 6
    assert result.maximum == 8
    assert result.score_range == 2
    assert result.mode == 8
    assert result.exact_mode_agreement == pytest.approx(0.6)


def test_score_statistics_reports_tied_modes_and_small_samples():
    tied = _score_statistics([6, 6, 8, 8])
    assert tied.mode == "6, 8"
    assert tied.exact_mode_agreement == pytest.approx(0.5)

    singleton = _score_statistics([7])
    assert singleton.sample_standard_deviation is None
    assert singleton.mode == 7

    empty = _score_statistics([])
    assert empty.valid_repeats == 0
    assert empty.mean is None
    assert empty.exact_mode_agreement is None


def test_reproducibility_workbook_has_runs_and_completed_only_statistics(
    tmp_path: Path, monkeypatch
):
    monkeypatch.setenv("GROUND_TRUTH_SCORE_FINDER_MODEL", "ground-truth-model")
    configs = [
        {"key": "analyst_alpha", "resolved_model_name": "alpha-model"},
        {"key": "analyst_beta", "resolved_model_name": "beta-model"},
    ]
    records = [
        ReproducibilityRunRecord(
            repeat=1,
            repeat_directory="case_1",
            status="completed",
            pdf_filename="case.pdf",
            analyst_scores={"analyst_alpha": 6, "analyst_beta": 7},
            masked_rucam_score=8,
            masked_rucam_category="Probable",
        ),
        ReproducibilityRunRecord(
            repeat=2,
            repeat_directory="case_2",
            status="completed",
            pdf_filename="case.pdf",
            analyst_scores={"analyst_alpha": 8, "analyst_beta": 7},
            masked_rucam_score=9,
            masked_rucam_category="Highly probable",
        ),
        ReproducibilityRunRecord(
            repeat=3,
            repeat_directory="case_3",
            status="failed",
            pdf_filename="case.pdf",
            analyst_scores={"analyst_alpha": 10, "analyst_beta": None},
            error="ProviderRequestError (status_code=502)",
        ),
    ]
    output_path = tmp_path / "case" / "summary.xlsx"

    written = _write_reproducibility_workbook(
        records=records,
        enabled_analyst_configs=configs,
        enable_score_masking=True,
        output_path=output_path,
    )

    assert written == output_path
    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == ["Runs", "Reproducibility"]
    runs = list(workbook["Runs"].values)
    assert runs[0] == (
        "repeat",
        "repeat_directory",
        "status",
        "pdf_filename",
        "analyst_alpha_score",
        "analyst_beta_score",
        "masked_rucam_score",
        "masked_rucam_category",
        "error",
    )
    assert runs[3][2] == "failed"
    assert runs[3][4] == 10

    aggregate_rows = list(workbook["Reproducibility"].values)
    aggregate = {row[0]: row for row in aggregate_rows[1:]}
    alpha = aggregate["analyst_alpha"]
    assert alpha[1] == "alpha-model"
    assert alpha[2] == 2
    assert alpha[3] == 7
    assert alpha[5:8] == (6, 8, 2)
    assert aggregate["ground_truth_rucam"][1] == "ground-truth-model"
    assert aggregate["ground_truth_rucam"][2] == 2


def test_reproducibility_runs_sorted_pdfs_and_repeats_in_isolated_directories(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    for name in ("case-b.pdf", "case-a.pdf"):
        (input_dir / name).write_bytes(b"%PDF-1.4")
    calls = []

    def fake_run_pdf_analysis(**kwargs):
        pdf_path = kwargs["pdf_path"]
        result_dir = kwargs["pdf_output_dir"]
        context = kwargs["run_context"]
        calls.append(
            (
                pdf_path.name,
                result_dir.relative_to(output_dir).as_posix(),
                context.repeat_index,
                kwargs["enable_score_masking"],
                kwargs["strict_scoring"],
                kwargs["max_restarts"],
            )
        )
        return PdfRunResult(
            row={
                "pdf_filename": pdf_path.name,
                "analyst_alpha": 6,
                "analyst_beta": 7,
                "analyst_gamma": 8,
                "masked_rucam_score": None,
                "masked_rucam_category": None,
            },
            reused=False,
        )

    monkeypatch.setattr(
        "dili_rucam_agents.reproducibility._run_pdf_analysis",
        fake_run_pdf_analysis,
    )

    summaries = run_reproducibility_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
        repeats=2,
        strict_scoring=True,
        max_restarts=1,
    )

    assert [path.relative_to(output_dir).as_posix() for path in summaries] == [
        "case-a/summary.xlsx",
        "case-b/summary.xlsx",
    ]
    assert calls == [
        ("case-a.pdf", "case-a/case-a_1", 1, False, True, 1),
        ("case-a.pdf", "case-a/case-a_2", 2, False, True, 1),
        ("case-b.pdf", "case-b/case-b_1", 1, False, True, 1),
        ("case-b.pdf", "case-b/case-b_2", 2, False, True, 1),
    ]
    assert not list(output_dir.rglob("batch_summary.xlsx"))


@pytest.mark.parametrize("value", (0, -1, True, 1.5, "5"))
def test_validate_repeats_rejects_non_positive_integers(value):
    with pytest.raises(ValueError, match="positive integer"):
        validate_repeats(value)


def test_reproducibility_empty_input_returns_no_summaries(tmp_path: Path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    assert (
        run_reproducibility_folder(
            input_dir=str(input_dir), output_dir=str(tmp_path / "output")
        )
        == []
    )


def test_reproducibility_persists_failed_repeat_then_stops(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    for name in ("case-a.pdf", "case-b.pdf"):
        (input_dir / name).write_bytes(b"%PDF-1.4")
    calls = []

    def fake_run_pdf_analysis(**kwargs):
        pdf_path = kwargs["pdf_path"]
        repeat = kwargs["run_context"].repeat_index
        calls.append((pdf_path.name, repeat))
        row = {
            "pdf_filename": pdf_path.name,
            "analyst_alpha": 6,
            "analyst_beta": None,
            "analyst_gamma": None,
            "masked_rucam_score": None,
            "masked_rucam_category": None,
        }
        if repeat == 2:
            raise PdfRunFailure(
                pdf_path=pdf_path,
                row=row,
                diagnostic="ProviderRequestError (status_code=502)",
            )
        return PdfRunResult(row=row, reused=False)

    monkeypatch.setattr(
        "dili_rucam_agents.reproducibility._run_pdf_analysis",
        fake_run_pdf_analysis,
    )

    with pytest.raises(RuntimeError, match="case-a.pdf repeat 2"):
        run_reproducibility_folder(
            input_dir=str(input_dir), output_dir=str(output_dir), repeats=3
        )

    assert calls == [("case-a.pdf", 1), ("case-a.pdf", 2)]
    workbook = load_workbook(output_dir / "case-a" / "summary.xlsx", data_only=True)
    runs = list(workbook["Runs"].values)
    assert [row[2] for row in runs[1:]] == ["completed", "failed"]
    aggregate = list(workbook["Reproducibility"].values)
    alpha = next(row for row in aggregate[1:] if row[0] == "analyst_alpha")
    assert alpha[2] == 1


def test_reproducibility_skips_only_compatible_completed_repeats(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    for repeat in (1, 2):
        repeat_dir = output_dir / "case" / f"case_{repeat}"
        _write_complete_reports(repeat_dir)
        (repeat_dir / "run_status.json").write_text(
            json.dumps(
                {
                    "pdf_filename": "case.pdf",
                    "status": "completed",
                    "masking_enabled": False,
                    "strict_scoring": False,
                    "enabled_analysts": [
                        "analyst_alpha",
                        "analyst_beta",
                        "analyst_gamma",
                    ],
                    "mode": "reproducibility",
                    "repeat_index": repeat,
                }
            ),
            encoding="utf-8",
        )
    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete",
        lambda *args, **kwargs: True,
    )
    guard = Mock(side_effect=AssertionError("completed repeat reran"))
    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", guard)

    summaries = run_reproducibility_folder(
        input_dir=str(input_dir), output_dir=str(output_dir), repeats=2
    )

    guard.assert_not_called()
    assert summaries == [output_dir / "case" / "summary.xlsx"]
    workbook = load_workbook(summaries[0], data_only=True)
    assert [row[0] for row in list(workbook["Runs"].values)[1:]] == [1, 2]


def test_reproducibility_force_rerun_disables_skip_and_resume(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    completion_guard = Mock(side_effect=AssertionError("force checked completion"))
    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete", completion_guard
    )
    resumes = []

    def fake_run_end_to_end(pdf_path, output_dir=None, **kwargs):
        resumes.append(kwargs["resume"])
        _write_complete_reports(Path(output_dir))
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    run_reproducibility_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
        repeats=2,
        force_rerun=True,
    )

    completion_guard.assert_not_called()
    assert resumes == [False, False]


def test_reproducibility_expansion_runs_only_the_new_repeat(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    for repeat in (1, 2):
        repeat_dir = output_dir / "case" / f"case_{repeat}"
        _write_complete_reports(repeat_dir)
        (repeat_dir / "run_status.json").write_text(
            json.dumps(
                {
                    "pdf_filename": "case.pdf",
                    "status": "completed",
                    "masking_enabled": False,
                    "strict_scoring": False,
                    "mode": "reproducibility",
                    "repeat_index": repeat,
                }
            ),
            encoding="utf-8",
        )
    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete",
        lambda pdf_path, output_dir, **kwargs: Path(output_dir).name
        in {"case_1", "case_2"},
    )
    executed = []

    def fake_run_end_to_end(pdf_path, output_dir=None, **kwargs):
        executed.append(Path(output_dir).name)
        _write_complete_reports(Path(output_dir))
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    run_reproducibility_folder(
        input_dir=str(input_dir), output_dir=str(output_dir), repeats=3
    )

    assert executed == ["case_3"]


def test_reproducibility_contraction_ignores_and_preserves_higher_directories(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    higher_dir = output_dir / "case" / "case_3"
    higher_dir.mkdir(parents=True)
    sentinel = higher_dir / "keep.txt"
    sentinel.write_text("keep", encoding="utf-8")

    def fake_run_pdf_analysis(**kwargs):
        pdf_path = kwargs["pdf_path"]
        return PdfRunResult(
            row={
                "pdf_filename": pdf_path.name,
                "analyst_alpha": 6,
                "analyst_beta": 6,
                "analyst_gamma": 6,
                "masked_rucam_score": None,
                "masked_rucam_category": None,
            },
            reused=False,
        )

    monkeypatch.setattr(
        "dili_rucam_agents.reproducibility._run_pdf_analysis",
        fake_run_pdf_analysis,
    )

    run_reproducibility_folder(
        input_dir=str(input_dir), output_dir=str(output_dir), repeats=2
    )

    assert sentinel.read_text(encoding="utf-8") == "keep"
    workbook = load_workbook(output_dir / "case" / "summary.xlsx", data_only=True)
    assert [row[0] for row in list(workbook["Runs"].values)[1:]] == [1, 2]


def test_reproducibility_failure_sinks_keep_raw_provider_text_private(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    secret = "sk-live-REPRODUCIBILITY-SECRET"
    clinical_text = "Patient Jane Doe ALT 980 after Drug Q"

    class ProviderRequestError(RuntimeError):
        status_code = 502

    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete",
        lambda *args, **kwargs: False,
    )

    def fail_run(*args, **kwargs):
        raise ProviderRequestError(
            f"Authorization: Bearer {secret}; request body: {clinical_text}"
        )

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fail_run)

    with pytest.raises(RuntimeError) as exc_info:
        run_reproducibility_folder(
            input_dir=str(input_dir),
            output_dir=str(output_dir),
            repeats=1,
            debug=True,
        )

    repeat_dir = output_dir / "case" / "case_1"
    status_text = (repeat_dir / "run_status.json").read_text()
    log_text = (repeat_dir / "case.log").read_text()
    workbook = load_workbook(output_dir / "case" / "summary.xlsx", data_only=True)
    workbook_text = " ".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    diagnostics = (str(exc_info.value), status_text, log_text, workbook_text)
    for diagnostic in diagnostics:
        assert secret not in diagnostic
        assert clinical_text not in diagnostic
        assert "Authorization" not in diagnostic
        assert "request body" not in diagnostic
        assert "ProviderRequestError" in diagnostic
        assert "status_code=502" in diagnostic


def test_reproducibility_success_has_no_log_without_debug(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete",
        lambda *args, **kwargs: False,
    )

    def fake_run_end_to_end(pdf_path, output_dir=None, **kwargs):
        _write_complete_reports(Path(output_dir))
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    run_reproducibility_folder(
        input_dir=str(input_dir), output_dir=str(output_dir), repeats=1
    )

    assert not (output_dir / "case" / "case_1" / "case.log").exists()


def test_batch_cli_dispatches_reproducibility_with_default_five(
    tmp_path: Path, monkeypatch, capsys
):
    captured = {}

    def fake_run_reproducibility_folder(**kwargs):
        captured.update(kwargs)
        return [Path(kwargs["output_dir"]) / "case" / "summary.xlsx"]

    monkeypatch.setattr(
        "dili_rucam_agents.reproducibility.run_reproducibility_folder",
        fake_run_reproducibility_folder,
    )

    _main(
        [
            str(tmp_path / "input"),
            str(tmp_path / "output"),
            "--reproducibility",
            "--mask-scores",
            "--analyst-delta",
            "--analyst-restarts",
            "1",
        ]
    )

    assert captured["repeats"] == 5
    assert captured["enable_score_masking"] is True
    assert captured["use_analyst_delta"] is True
    assert captured["max_restarts"] == 1
    assert "summary.xlsx" in capsys.readouterr().out


def test_batch_cli_forwards_explicit_repeat_count(tmp_path: Path, monkeypatch):
    captured = {}
    monkeypatch.setattr(
        "dili_rucam_agents.reproducibility.run_reproducibility_folder",
        lambda **kwargs: captured.update(kwargs) or [],
    )
    _main(
        [
            str(tmp_path / "input"),
            str(tmp_path / "output"),
            "--reproducibility",
            "--repeats",
            "7",
        ]
    )
    assert captured["repeats"] == 7


def test_batch_cli_rejects_repeats_without_reproducibility(tmp_path: Path, capsys):
    with pytest.raises(SystemExit, match="2"):
        _main(
            [
                str(tmp_path / "input"),
                str(tmp_path / "output"),
                "--repeats",
                "3",
            ]
        )
    assert "--repeats requires --reproducibility" in capsys.readouterr().err


@pytest.mark.parametrize("value", ("0", "-1", "not-an-integer"))
def test_batch_cli_rejects_invalid_repeat_values(tmp_path: Path, value: str):
    with pytest.raises(SystemExit, match="2"):
        _main(
            [
                str(tmp_path / "input"),
                str(tmp_path / "output"),
                "--reproducibility",
                "--repeats",
                value,
            ]
        )


def test_batch_cli_keeps_regular_dispatch_unchanged(tmp_path: Path, monkeypatch):
    captured = {}

    def fake_run_batch_folder(**kwargs):
        captured.update(kwargs)
        return Path(kwargs["output_dir"]) / "batch_summary.xlsx"

    monkeypatch.setattr(
        "dili_rucam_agents.batch.run_batch_folder", fake_run_batch_folder
    )

    _main([str(tmp_path / "input"), str(tmp_path / "output")])

    assert captured == {
        "input_dir": str(tmp_path / "input"),
        "output_dir": str(tmp_path / "output"),
        "prompt_path": None,
        "enable_score_masking": False,
        "strict_scoring": False,
        "use_analyst_delta": False,
        "use_analyst_epsilon": False,
        "use_analyst_zeta": False,
        "use_analyst_eta": False,
        "debug": False,
        "force_rerun": False,
        "max_restarts": 2,
    }
