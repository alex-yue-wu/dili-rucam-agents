import json
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook
import pytest

import dili_rucam_agents.pipeline as pipeline_module
from dili_rucam_agents.batch import (
    _is_pdf_run_complete,
    extract_ground_truth_rucam_category,
    extract_ground_truth_rucam_score,
    extract_section_c_json,
    run_batch_folder,
)
from dili_rucam_agents.crew.crew import AnalystAttemptEvent, AnalystExecutionError
from dili_rucam_agents.pipeline import is_end_to_end_complete, run_end_to_end
from dili_rucam_agents.validators.analyst_report import (
    AnalystReportValidationError,
    validate_analyst_report,
)


def complete_report(total_score: int = 6) -> str:
    rucam_scores = {
        "time_to_onset": 2,
        "course": 1,
        "risk_factors": 0,
        "concomitant_drugs": 0,
        "other_causes_excluded": 2,
        "known_hepatotoxicity": 1,
        "rechallenge": 0,
    }
    if total_score == 7:
        rucam_scores["known_hepatotoxicity"] = 2
    elif total_score == 8:
        rucam_scores["course"] = 2
        rucam_scores["known_hepatotoxicity"] = 2
    elif total_score != 6:
        raise ValueError("complete_report only supports scores from 6 through 8")
    payload = {
        "injury_pattern": "hepatocellular",
        "R_ratio": 6.4,
        "rucam_scores": rucam_scores,
        "total_score": total_score,
        "category": "Probable",
    }
    return (
        "## SECTION A\n\nClinical summary\n\n"
        f"## SECTION B\n\n| Item | Score |\n| --- | --- |\n| Total | {total_score} |\n\n"
        f"## SECTION C\n\n```json\n{json.dumps(payload)}\n```\n"
    )


def write_complete_reports(result_dir: Path) -> None:
    result_dir.mkdir(parents=True, exist_ok=True)
    for report_name in (
        "analyst-alpha_report.md",
        "analyst-beta_report.md",
        "analyst-gamma_report.md",
    ):
        (result_dir / report_name).write_text(complete_report(), encoding="utf-8")


def test_extract_section_c_json_parses_last_json_block():
    report_text = """
Intro

```json
{"ignored": true}
```

## SECTION C

```json
{"total_score": 8, "category": "Probable"}
```
"""

    payload = extract_section_c_json(report_text)

    assert payload["total_score"] == 8
    assert payload["category"] == "Probable"


def test_extract_section_c_json_recovers_common_malformed_llm_json():
    report_text = """
## SECTION C

```json
{
{
  "injury_pattern": "hepatocellular",",
  "R_ratio": 6.44,,
  "rucam_scores": { {
    "time_to_onset": 2,,
    "course": 0,,
    "risk_factors": 0,,
    "concomitant_drugs": 0,,
    "other_causes_excluded": 0,,
    "known_hepatotoxicity": 2,,
    "rechallenge": 0
  }, },
  "total_score": 4,,
  "category": "Possible",",
  "rules_version": "RUCAM_ULN_LiverTox_v1",",
  "notes": [ [
    "ULN inferred using LiverTox defaults",",
    "Course inferred from qualitative description (no numeric trajectory reported)",)
  ] ]
} ]
}
```
"""

    payload = extract_section_c_json(report_text)

    assert payload["injury_pattern"] == "hepatocellular"
    assert payload["R_ratio"] == 6.44
    assert payload["total_score"] == 4
    assert payload["category"] == "Possible"
    assert payload["rucam_scores"]["time_to_onset"] == 2
    assert payload["notes"] == [
        "ULN inferred using LiverTox defaults",
        "Course inferred from qualitative description (no numeric trajectory reported",
    ]

    strict_report = "## SECTION A\n\nSummary\n\n## SECTION B\n\nTable\n\n" + report_text
    with pytest.raises(AnalystReportValidationError, match="Invalid SECTION C JSON"):
        validate_analyst_report(strict_report)


def test_extract_section_c_json_parses_unfenced_json_below_section_c():
    report_text = """
## SECTION C — MACHINE-READABLE JSON

{
  "injury_pattern": "Not reported",
  "R_ratio": null,
  "rucam_scores": {
    "time_to_onset": 0,
    "course": 0,
    "risk_factors": 0,
    "concomitant_drugs": 0,
    "other_causes_excluded": -2,
    "known_hepatotoxicity": 0,
    "rechallenge": 0
  },
  "total_score": -2,
  "category": "Excluded",
  "notes": [
    "No extractable case content was provided in the case bundle."
  ],
  "rules_version": "RUCAM_ULN_LiverTox_v1"
}
"""

    payload = extract_section_c_json(report_text)

    assert payload["total_score"] == -2
    assert payload["category"] == "Excluded"
    assert payload["rucam_scores"]["other_causes_excluded"] == -2


def test_extract_section_c_json_raises_helpful_error_for_summary_placeholder():
    report_text = "See complete Sections A, B, and C above."

    try:
        extract_section_c_json(report_text)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "summary placeholder" in str(exc)


def test_extract_section_c_json_raises_helpful_error_for_truncated_report():
    report_text = """
## SECTION A

Complete narrative.

## SECTION B

| RUCAM Item | Score |
| --- | --- |
| Total | 6 |
"""

    try:
        extract_section_c_json(report_text)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "truncated before SECTION C" in str(exc)


def test_extract_ground_truth_rucam_score_reads_stable_field():
    report = """
# Ground Truth RUCAM Score Report

## Stable Fields
```text
GROUND_TRUTH_RUCAM_SCORE: 8
```
"""

    assert extract_ground_truth_rucam_score(report) == 8


def test_extract_ground_truth_rucam_category_reads_stable_field():
    report = """
# Ground Truth RUCAM Score Report

## Stable Fields
```text
GROUND_TRUTH_RUCAM_SCORE: 8
GROUND_TRUTH_RUCAM_CATEGORY: Probable
```
"""

    assert extract_ground_truth_rucam_category(report) == "Probable"


def test_run_batch_folder_writes_summary_workbook(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")
    (input_dir / "case-b.pdf").write_bytes(b"%PDF-1.4")

    def fake_run_end_to_end(
        pdf_path: str,
        prompt_path: str | None = None,
        output_dir: str | None = None,
        **kwargs,
    ) -> str:
        result_dir = Path(output_dir)
        result_dir.mkdir(parents=True, exist_ok=True)
        for report_name in (
            "analyst-alpha_report.md",
            "analyst-beta_report.md",
            "analyst-gamma_report.md",
        ):
            result_dir.joinpath(report_name).write_text(
                complete_report(7),
                encoding="utf-8",
            )
        if kwargs.get("enable_score_masking"):
            result_dir.joinpath("ground-truth-rucam-score_report.md").write_text(
                "# Ground Truth RUCAM Score Report\n\n"
                "## Stable Fields\n"
                "```text\nGROUND_TRUTH_RUCAM_SCORE: 8\nGROUND_TRUTH_RUCAM_CATEGORY: Probable\n```\n",
                encoding="utf-8",
            )
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    summary_path = run_batch_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
        enable_score_masking=True,
    )

    workbook = load_workbook(summary_path)
    sheet = workbook["Batch Summary"]
    headers = [cell.value for cell in sheet[1]]
    first_row = [cell.value for cell in sheet[2]]
    row_map = dict(zip(headers, first_row))
    log_path = output_dir / "case-a" / "case-a.log"

    assert summary_path.name == "batch_summary.xlsx"
    assert headers == [
        "pdf_filename",
        "gpt-5.5",
        "gemini-3.1-pro-preview",
        "moonshotai/kimi-k2.5",
        "masked_rucam_score",
        "masked_rucam_category",
    ]
    assert row_map["pdf_filename"] == "case-a.pdf"
    assert row_map["masked_rucam_score"] == 8
    assert row_map["masked_rucam_category"] == "Probable"
    assert row_map["gpt-5.5"] == 7
    assert row_map["gemini-3.1-pro-preview"] == 7
    assert row_map["moonshotai/kimi-k2.5"] == 7
    assert not log_path.exists()


def test_run_batch_folder_forwards_restart_limit_and_resume(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    captured_kwargs = {}

    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete", lambda *args, **kwargs: False
    )

    def fake_run_end_to_end(pdf_path, output_dir=None, **kwargs):
        captured_kwargs.update(kwargs)
        write_complete_reports(Path(output_dir))
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)
    run_batch_folder(
        input_dir=str(input_dir), output_dir=str(output_dir), max_restarts=1
    )
    assert captured_kwargs["max_restarts"] == 1
    assert captured_kwargs["resume"] is True


def test_force_rerun_disables_pipeline_resume(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")
    captured_kwargs = {}

    def fake_run_end_to_end(pdf_path, output_dir=None, **kwargs):
        captured_kwargs.update(kwargs)
        write_complete_reports(Path(output_dir))
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)
    run_batch_folder(
        input_dir=str(input_dir), output_dir=str(output_dir), force_rerun=True
    )
    assert captured_kwargs["resume"] is False


def test_rerun_skips_completed_pdf_then_runs_failed_pdf(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    for name in ("case-a.pdf", "case-b.pdf"):
        (input_dir / name).write_bytes(b"%PDF-1.4")
    case_a_dir = output_dir / "case-a"
    write_complete_reports(case_a_dir)
    (case_a_dir / "run_status.json").write_text(
        json.dumps(
            {
                "pdf_filename": "case-a.pdf",
                "status": "completed",
                "masking_enabled": False,
                "strict_scoring": False,
                "enabled_analysts": [
                    "analyst_alpha",
                    "analyst_beta",
                    "analyst_gamma",
                ],
            }
        ),
        encoding="utf-8",
    )
    calls = []

    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete",
        lambda pdf_path, *args, **kwargs: Path(pdf_path).name == "case-a.pdf",
    )

    def fake_run_end_to_end(pdf_path, output_dir=None, **kwargs):
        calls.append(Path(pdf_path).name)
        write_complete_reports(Path(output_dir))
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)
    summary_path = run_batch_folder(
        input_dir=str(input_dir), output_dir=str(output_dir)
    )
    workbook = load_workbook(summary_path)
    summary_workbook_rows = [
        workbook.active.cell(row=index, column=1).value for index in (2, 3)
    ]
    assert calls == ["case-b.pdf"]
    assert summary_workbook_rows == ["case-a.pdf", "case-b.pdf"]


def test_run_batch_folder_rejects_more_than_two_restarts(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(
        "dili_rucam_agents.batch.run_end_to_end",
        lambda *args, **kwargs: calls.append("called"),
    )
    with pytest.raises(ValueError, match="0 through 2"):
        run_batch_folder(
            input_dir=str(tmp_path / "input"),
            output_dir=str(tmp_path / "output"),
            max_restarts=3,
        )
    assert calls == []


def test_run_batch_folder_records_analyst_failure_details(tmp_path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case.pdf").write_bytes(b"%PDF-1.4")

    def fake_run_end_to_end(*args, **kwargs):
        raise AnalystExecutionError(
            analyst_key="analyst_beta",
            attempts=3,
            failure_kind="validation",
            last_error="Invalid SECTION C JSON",
        )

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    with pytest.raises(RuntimeError, match="case.pdf"):
        run_batch_folder(input_dir=str(input_dir), output_dir=str(output_dir))

    status = json.loads(
        (output_dir / "case" / "run_status.json").read_text(encoding="utf-8")
    )
    assert status["failed_analyst"] == "analyst_beta"
    assert status["attempts"] == 3
    assert status["failure_kind"] == "validation"


def test_run_batch_folder_writes_log_only_in_debug_mode(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")

    def fake_run_end_to_end(
        pdf_path: str,
        prompt_path: str | None = None,
        output_dir: str | None = None,
        **kwargs,
    ) -> str:
        result_dir = Path(output_dir)
        result_dir.mkdir(parents=True, exist_ok=True)
        for report_name in (
            "analyst-alpha_report.md",
            "analyst-beta_report.md",
            "analyst-gamma_report.md",
        ):
            result_dir.joinpath(report_name).write_text(
                complete_report(7),
                encoding="utf-8",
            )
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    run_batch_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
        debug=True,
    )

    log_path = output_dir / "case-a" / "case-a.log"
    assert log_path.exists()
    assert "Running PDF:" in log_path.read_text(encoding="utf-8")


def test_run_end_to_end_forwards_strict_scoring(monkeypatch):
    captured_kwargs = {}

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        captured_kwargs.update(kwargs)
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", fake_run_crew)

    result = run_end_to_end("example.pdf", strict_scoring=True)

    assert result == "ok"
    assert captured_kwargs["strict_scoring"] is True


def test_run_end_to_end_persists_invalid_attempt_then_valid_report(
    tmp_path, monkeypatch
):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        callback = kwargs["on_attempt"]
        callback(AnalystAttemptEvent("analyst_alpha", 1, 3, "running"))
        callback(
            AnalystAttemptEvent(
                "analyst_alpha",
                1,
                3,
                "validation_failed",
                error="Missing SECTION C",
                report_text="incomplete",
            )
        )
        callback(AnalystAttemptEvent("analyst_alpha", 2, 3, "running"))
        callback(
            AnalystAttemptEvent(
                "analyst_alpha", 2, 3, "completed", report_text=complete_report()
            )
        )
        return "ok", {"analyst_alpha": complete_report()}

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", fake_run_crew)
    run_end_to_end(str(pdf_path), output_dir=str(output_dir))

    attempt_paths = list((output_dir / "attempts").glob("analyst-alpha_*.invalid.md"))
    assert len(attempt_paths) == 1
    assert attempt_paths[0].read_text() == "incomplete"
    assert (
        validate_analyst_report(
            (output_dir / "analyst-alpha_report.md").read_text()
        ).payload.total_score
        == 6
    )


def test_invalid_attempt_artifacts_and_history_accumulate_across_invocations(
    tmp_path, monkeypatch
):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"
    invalid_outputs = iter(("first invalid output", "second invalid output"))

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        report_text = next(invalid_outputs)
        callback = kwargs["on_attempt"]
        callback(AnalystAttemptEvent("analyst_alpha", 1, 1, "running"))
        callback(
            AnalystAttemptEvent(
                "analyst_alpha",
                1,
                1,
                "validation_failed",
                error="Missing SECTION C",
                report_text=report_text,
            )
        )
        raise AnalystExecutionError(
            analyst_key="analyst_alpha",
            attempts=1,
            failure_kind="validation",
            last_error="Missing SECTION C",
        )

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", fake_run_crew)

    for _ in range(2):
        with pytest.raises(AnalystExecutionError):
            run_end_to_end(
                str(pdf_path),
                output_dir=str(output_dir),
                max_restarts=0,
            )

    attempt_paths = sorted((output_dir / "attempts").glob("*.invalid.md"))
    assert len(attempt_paths) == 2
    assert attempt_paths[0] != attempt_paths[1]
    assert {path.read_text() for path in attempt_paths} == {
        "first invalid output",
        "second invalid output",
    }
    manifest = json.loads(
        (output_dir / "analyst_checkpoints.json").read_text(encoding="utf-8")
    )
    entry = manifest["analysts"]["analyst_alpha"]
    assert entry["total_attempts"] == 2
    assert [item["total_attempt"] for item in entry["attempt_history"]] == [1, 2]
    assert all(item["artifact_file"] for item in entry["attempt_history"])


def test_run_end_to_end_resume_false_supplies_no_completed_reports(
    tmp_path, monkeypatch
):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"
    output_dir.mkdir()
    (output_dir / "analyst-alpha_report.md").write_text(
        complete_report(), encoding="utf-8"
    )
    captured_kwargs = {}

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        captured_kwargs.update(kwargs)
        return "ok", {}

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", fake_run_crew)
    run_end_to_end(str(pdf_path), output_dir=str(output_dir), resume=False)

    assert captured_kwargs["completed_reports"] == {}


def test_run_end_to_end_forwards_max_restarts(monkeypatch):
    captured_kwargs = {}

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        captured_kwargs.update(kwargs)
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", fake_run_crew)
    run_end_to_end("example.pdf", max_restarts=1)

    assert captured_kwargs["max_restarts"] == 1


def test_run_end_to_end_passes_fingerprinted_instruction_contracts_unchanged(
    tmp_path, monkeypatch
):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"
    checkpoint_contexts = []
    captured_kwargs = {}
    original_builder = pipeline_module._build_checkpoint_context

    def capture_checkpoint_context(**kwargs):
        context = original_builder(**kwargs)
        checkpoint_contexts.append(context)
        return context

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        captured_kwargs.update(kwargs)
        return "ok", dict(kwargs["completed_reports"])

    monkeypatch.setattr(
        pipeline_module, "_build_checkpoint_context", capture_checkpoint_context
    )
    monkeypatch.setattr(pipeline_module, "run_crew", fake_run_crew)

    run_end_to_end(str(pdf_path), output_dir=str(output_dir))

    assert len(checkpoint_contexts) == 1
    assert (
        captured_kwargs["instruction_contracts"]
        is checkpoint_contexts[0].instruction_contracts
    )


def test_run_end_to_end_omits_attempt_handler_without_output_dir(monkeypatch):
    captured_kwargs = {}

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        captured_kwargs.update(kwargs)
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", fake_run_crew)
    run_end_to_end("example.pdf")

    assert captured_kwargs["on_attempt"] is None


def test_later_failure_preserves_completed_alpha_checkpoint(tmp_path, monkeypatch):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"

    def fake_run_crew(pdf_path, prompt_path=None, **kwargs):
        callback = kwargs["on_attempt"]
        callback(AnalystAttemptEvent("analyst_alpha", 1, 3, "running"))
        callback(
            AnalystAttemptEvent(
                "analyst_alpha", 1, 3, "completed", report_text=complete_report()
            )
        )
        callback(AnalystAttemptEvent("analyst_beta", 1, 3, "running"))
        callback(
            AnalystAttemptEvent(
                "analyst_beta",
                1,
                3,
                "validation_failed",
                error="Missing SECTION C",
                report_text="incomplete",
            )
        )
        raise AnalystExecutionError(
            analyst_key="analyst_beta",
            attempts=3,
            failure_kind="validation",
            last_error="Missing SECTION C",
        )

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", fake_run_crew)
    with pytest.raises(AnalystExecutionError):
        run_end_to_end(str(pdf_path), output_dir=str(output_dir))
    assert (
        validate_analyst_report(
            (output_dir / "analyst-alpha_report.md").read_text()
        ).payload.total_score
        == 6
    )


def test_second_invocation_supplies_all_checkpointed_reports(tmp_path, monkeypatch):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"
    analyst_keys = ("analyst_alpha", "analyst_beta", "analyst_gamma")

    def seed_run(pdf_path, prompt_path=None, **kwargs):
        for key in analyst_keys:
            kwargs["on_attempt"](AnalystAttemptEvent(key, 1, 3, "running"))
            kwargs["on_attempt"](
                AnalystAttemptEvent(
                    key, 1, 3, "completed", report_text=complete_report()
                )
            )
        return "ok", {key: complete_report() for key in analyst_keys}

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", seed_run)
    run_end_to_end(str(pdf_path), output_dir=str(output_dir))
    captured_kwargs = {}

    def resume_run(pdf_path, prompt_path=None, **kwargs):
        captured_kwargs.update(kwargs)
        return "ok", dict(kwargs["completed_reports"])

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", resume_run)
    run_end_to_end(str(pdf_path), output_dir=str(output_dir))
    assert set(captured_kwargs["completed_reports"]) == set(analyst_keys)


def test_is_end_to_end_complete_uses_validated_checkpoint_manifest(
    tmp_path, monkeypatch
):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"
    analyst_keys = ("analyst_alpha", "analyst_beta", "analyst_gamma")

    def seed_run(pdf_path, prompt_path=None, **kwargs):
        for key in analyst_keys:
            kwargs["on_attempt"](AnalystAttemptEvent(key, 1, 3, "running"))
            kwargs["on_attempt"](
                AnalystAttemptEvent(
                    key, 1, 3, "completed", report_text=complete_report()
                )
            )
        return "ok", {key: complete_report() for key in analyst_keys}

    monkeypatch.setattr("dili_rucam_agents.pipeline.run_crew", seed_run)
    run_end_to_end(str(pdf_path), output_dir=str(output_dir))
    manifest_path = output_dir / "analyst_checkpoints.json"
    manifest_before = manifest_path.read_text(encoding="utf-8")

    assert is_end_to_end_complete(str(pdf_path), str(output_dir))
    assert manifest_path.read_text(encoding="utf-8") == manifest_before


def test_is_end_to_end_complete_does_not_write_contracted_topology(
    tmp_path, monkeypatch
):
    pdf_path = tmp_path / "example.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    output_dir = tmp_path / "case"
    analyst_keys = (
        "analyst_alpha",
        "analyst_beta",
        "analyst_gamma",
        "analyst_delta",
    )

    def seed_run(pdf_path, prompt_path=None, **kwargs):
        for key in analyst_keys:
            kwargs["on_attempt"](AnalystAttemptEvent(key, 1, 3, "running"))
            kwargs["on_attempt"](
                AnalystAttemptEvent(
                    key, 1, 3, "completed", report_text=complete_report()
                )
            )
        return "ok", {key: complete_report() for key in analyst_keys}

    monkeypatch.setattr(pipeline_module, "run_crew", seed_run)
    run_end_to_end(
        str(pdf_path),
        output_dir=str(output_dir),
        use_analyst_delta=True,
    )
    manifest_path = output_dir / "analyst_checkpoints.json"
    manifest_before = manifest_path.read_bytes()
    writer = Mock(side_effect=AssertionError("completion check wrote manifest"))
    monkeypatch.setattr("dili_rucam_agents.checkpoints.atomic_write_json", writer)

    assert is_end_to_end_complete(str(pdf_path), str(output_dir))
    writer.assert_not_called()
    assert manifest_path.read_bytes() == manifest_before


def test_run_batch_folder_skips_completed_pdf(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")

    result_dir = output_dir / "case-a"
    result_dir.mkdir()
    result_dir.joinpath("run_status.json").write_text(
        json.dumps(
            {
                "pdf_filename": "case-a.pdf",
                "status": "completed",
                "masking_enabled": True,
                "strict_scoring": False,
                "enabled_analysts": ["analyst_alpha", "analyst_beta", "analyst_gamma"],
            }
        ),
        encoding="utf-8",
    )
    for report_name in (
        "analyst-alpha_report.md",
        "analyst-beta_report.md",
        "analyst-gamma_report.md",
    ):
        result_dir.joinpath(report_name).write_text(
            complete_report(7),
            encoding="utf-8",
        )
    result_dir.joinpath("ground-truth-rucam-score_report.md").write_text(
        "# Ground Truth RUCAM Score Report\n\n"
        "## Stable Fields\n"
        "```text\nGROUND_TRUTH_RUCAM_SCORE: 8\nGROUND_TRUTH_RUCAM_CATEGORY: Probable\n```\n",
        encoding="utf-8",
    )
    result_dir.joinpath("masked-case-bundle_report.md").write_text(
        "# Masked Case Bundle Report\n\n## Stable Fields\n```text\nMASKED_RUCAM_SCORES: 8\n```\n",
        encoding="utf-8",
    )

    calls: list[str] = []

    def fake_run_end_to_end(*args, **kwargs):
        calls.append("called")
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)
    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete", lambda *args, **kwargs: True
    )

    run_batch_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
        enable_score_masking=True,
    )

    assert calls == []


def test_batch_contracted_topology_skips_without_mutating_checkpoint(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    pdf_path = input_dir / "case-a.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    result_dir = output_dir / "case-a"
    analyst_keys = (
        "analyst_alpha",
        "analyst_beta",
        "analyst_gamma",
        "analyst_delta",
    )

    def seed_run(pdf_path, prompt_path=None, **kwargs):
        for key in analyst_keys:
            kwargs["on_attempt"](AnalystAttemptEvent(key, 1, 3, "running"))
            kwargs["on_attempt"](
                AnalystAttemptEvent(
                    key, 1, 3, "completed", report_text=complete_report()
                )
            )
        return "ok", {key: complete_report() for key in analyst_keys}

    monkeypatch.setattr(pipeline_module, "run_crew", seed_run)
    run_end_to_end(str(pdf_path), output_dir=str(result_dir), use_analyst_delta=True)
    result_dir.joinpath("run_status.json").write_text(
        json.dumps(
            {
                "pdf_filename": "case-a.pdf",
                "status": "completed",
                "masking_enabled": False,
                "strict_scoring": False,
                "enabled_analysts": list(analyst_keys),
            }
        ),
        encoding="utf-8",
    )
    manifest_path = result_dir / "analyst_checkpoints.json"
    manifest_before = manifest_path.read_bytes()

    monkeypatch.setattr(
        "dili_rucam_agents.batch.run_end_to_end",
        Mock(side_effect=AssertionError("contracted batch reran pipeline")),
    )

    run_batch_folder(input_dir=str(input_dir), output_dir=str(output_dir))

    assert manifest_path.read_bytes() == manifest_before


def test_batch_topology_expansion_reenters_pipeline(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    pdf_path = input_dir / "case-a.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    result_dir = output_dir / "case-a"

    def seed_run(pdf_path, prompt_path=None, **kwargs):
        for key in ("analyst_alpha", "analyst_beta", "analyst_gamma"):
            kwargs["on_attempt"](AnalystAttemptEvent(key, 1, 3, "running"))
            kwargs["on_attempt"](
                AnalystAttemptEvent(
                    key, 1, 3, "completed", report_text=complete_report()
                )
            )
        return "ok", {
            key: complete_report()
            for key in ("analyst_alpha", "analyst_beta", "analyst_gamma")
        }

    monkeypatch.setattr(pipeline_module, "run_crew", seed_run)
    run_end_to_end(str(pdf_path), output_dir=str(result_dir))
    result_dir.joinpath("run_status.json").write_text(
        json.dumps(
            {
                "pdf_filename": "case-a.pdf",
                "status": "completed",
                "masking_enabled": False,
                "strict_scoring": False,
                "enabled_analysts": [
                    "analyst_alpha",
                    "analyst_beta",
                    "analyst_gamma",
                ],
            }
        ),
        encoding="utf-8",
    )
    calls = []

    def expanded_run(pdf_path, *, output_dir, **kwargs):
        calls.append(Path(pdf_path).name)
        write_complete_reports(Path(output_dir))
        Path(output_dir, "analyst-delta_report.md").write_text(
            complete_report(), encoding="utf-8"
        )
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", expanded_run)

    run_batch_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
        use_analyst_delta=True,
    )

    assert calls == ["case-a.pdf"]


def test_run_batch_folder_force_rerun_ignores_completed_status(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")

    result_dir = output_dir / "case-a"
    result_dir.mkdir()
    result_dir.joinpath("run_status.json").write_text(
        json.dumps(
            {
                "pdf_filename": "case-a.pdf",
                "status": "completed",
                "masking_enabled": False,
                "strict_scoring": False,
                "enabled_analysts": ["analyst_alpha", "analyst_beta", "analyst_gamma"],
            }
        ),
        encoding="utf-8",
    )
    for report_name in (
        "analyst-alpha_report.md",
        "analyst-beta_report.md",
        "analyst-gamma_report.md",
    ):
        result_dir.joinpath(report_name).write_text(
            complete_report(7),
            encoding="utf-8",
        )

    calls: list[str] = []

    def fake_run_end_to_end(
        pdf_path: str,
        prompt_path: str | None = None,
        output_dir: str | None = None,
        **kwargs,
    ) -> str:
        calls.append(Path(pdf_path).name)
        result_dir = Path(output_dir)
        for report_name in (
            "analyst-alpha_report.md",
            "analyst-beta_report.md",
            "analyst-gamma_report.md",
        ):
            result_dir.joinpath(report_name).write_text(
                complete_report(8),
                encoding="utf-8",
            )
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    run_batch_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
        force_rerun=True,
    )

    assert calls == ["case-a.pdf"]


def test_is_pdf_run_complete_requires_status_and_parseable_reports(tmp_path: Path):
    result_dir = tmp_path / "case-a"
    result_dir.mkdir()
    pdf_path = tmp_path / "case-a.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    assert not _is_pdf_run_complete(
        pdf_path=pdf_path,
        pdf_output_dir=result_dir,
        prompt_path=None,
        enabled_analyst_configs=[
            {"key": "analyst_alpha"},
            {"key": "analyst_beta"},
            {"key": "analyst_gamma"},
        ],
        enable_score_masking=True,
        strict_scoring=False,
        use_analyst_delta=False,
        use_analyst_epsilon=False,
        use_analyst_zeta=False,
        use_analyst_eta=False,
    )


def test_completed_status_delegates_all_compatibility_arguments(tmp_path, monkeypatch):
    pdf_path = tmp_path / "case-a.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    result_dir = tmp_path / "case-a"
    result_dir.mkdir()
    enabled_analyst_configs = [
        {"key": "analyst_alpha"},
        {"key": "analyst_beta"},
        {"key": "analyst_gamma"},
        {"key": "analyst_delta"},
    ]
    result_dir.joinpath("run_status.json").write_text(
        json.dumps(
            {
                "pdf_filename": "case-a.pdf",
                "status": "completed",
                "masking_enabled": True,
                "strict_scoring": True,
                "enabled_analysts": [
                    "analyst_alpha",
                    "analyst_beta",
                    "analyst_gamma",
                    "analyst_delta",
                ],
            }
        ),
        encoding="utf-8",
    )
    captured_calls = []
    validator_outcomes = iter((True, False))

    def fake_is_end_to_end_complete(pdf_arg, output_arg, prompt_path=None, **kwargs):
        captured_calls.append((pdf_arg, output_arg, prompt_path, kwargs))
        return next(validator_outcomes)

    monkeypatch.setattr(
        "dili_rucam_agents.batch.is_end_to_end_complete",
        fake_is_end_to_end_complete,
    )

    settings = {
        "pdf_path": pdf_path,
        "pdf_output_dir": result_dir,
        "prompt_path": str(tmp_path / "prompt.md"),
        "enabled_analyst_configs": enabled_analyst_configs,
        "enable_score_masking": True,
        "strict_scoring": True,
        "use_analyst_delta": True,
        "use_analyst_epsilon": False,
        "use_analyst_zeta": False,
        "use_analyst_eta": False,
    }

    assert _is_pdf_run_complete(**settings) is True
    assert _is_pdf_run_complete(**settings) is False
    assert captured_calls == [
        (
            str(pdf_path),
            str(result_dir),
            str(tmp_path / "prompt.md"),
            {
                "enable_score_masking": True,
                "strict_scoring": True,
                "use_analyst_delta": True,
                "use_analyst_epsilon": False,
                "use_analyst_zeta": False,
                "use_analyst_eta": False,
            },
        ),
        (
            str(pdf_path),
            str(result_dir),
            str(tmp_path / "prompt.md"),
            {
                "enable_score_masking": True,
                "strict_scoring": True,
                "use_analyst_delta": True,
                "use_analyst_epsilon": False,
                "use_analyst_zeta": False,
                "use_analyst_eta": False,
            },
        ),
    ]


def test_run_batch_folder_stops_on_error_and_marks_failed(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")
    (input_dir / "case-b.pdf").write_bytes(b"%PDF-1.4")

    calls: list[str] = []

    def fake_run_end_to_end(pdf_path: str, **kwargs) -> str:
        calls.append(Path(pdf_path).name)
        if Path(pdf_path).name == "case-a.pdf":
            raise RuntimeError("simulated failure")
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    try:
        run_batch_folder(
            input_dir=str(input_dir),
            output_dir=str(output_dir),
        )
        assert False, "expected RuntimeError"
    except RuntimeError as exc:
        assert "case-a.pdf" in str(exc)

    assert calls == ["case-a.pdf"]
    status = json.loads(
        (output_dir / "case-a" / "run_status.json").read_text(encoding="utf-8")
    )
    assert status["status"] == "failed"
    assert "RuntimeError" in status["error"]
    assert "simulated failure" not in status["error"]
    assert not (output_dir / "case-b" / "run_status.json").exists()


def test_batch_execution_exception_diagnostics_never_persist_raw_text(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")
    secret = "sk-live-TOKEN-DO-NOT-STORE"
    clinical_text = "Patient Jane Doe ALT 980 after Drug Q"

    class ProviderRequestError(RuntimeError):
        status_code = 502

    def fail_run(*args, **kwargs):
        raise ProviderRequestError(
            f"Authorization: Bearer {secret}; request body: {clinical_text}"
        )

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fail_run)

    with pytest.raises(RuntimeError) as exc_info:
        run_batch_folder(
            input_dir=str(input_dir),
            output_dir=str(output_dir),
            debug=True,
        )

    result_dir = output_dir / "case-a"
    status_text = result_dir.joinpath("run_status.json").read_text(encoding="utf-8")
    log_text = result_dir.joinpath("case-a.log").read_text(encoding="utf-8")
    workbook = load_workbook(output_dir / "batch_summary.xlsx", data_only=True)
    workbook_text = " ".join(
        str(cell.value)
        for row in workbook.active.iter_rows()
        for cell in row
        if cell.value is not None
    )
    diagnostics = (str(exc_info.value), status_text, log_text, workbook_text)
    for diagnostic in diagnostics:
        assert secret not in diagnostic
        assert clinical_text not in diagnostic
        assert "Authorization" not in diagnostic
        assert "request body" not in diagnostic
        assert "ProviderRequestError" in diagnostic
        assert "status_code=502" in diagnostic


def test_run_batch_folder_identifies_analyst_report_when_section_c_is_missing(
    tmp_path: Path, monkeypatch
):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")

    def fake_run_end_to_end(
        pdf_path: str,
        prompt_path: str | None = None,
        output_dir: str | None = None,
        **kwargs,
    ) -> str:
        result_dir = Path(output_dir)
        result_dir.mkdir(parents=True, exist_ok=True)
        result_dir.joinpath("analyst-alpha_report.md").write_text(
            complete_report(7),
            encoding="utf-8",
        )
        result_dir.joinpath("analyst-beta_report.md").write_text(
            "See complete Sections A, B, and C above.\n",
            encoding="utf-8",
        )
        result_dir.joinpath("analyst-gamma_report.md").write_text(
            complete_report(7),
            encoding="utf-8",
        )
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    try:
        run_batch_folder(
            input_dir=str(input_dir),
            output_dir=str(output_dir),
        )
        assert False, "expected RuntimeError"
    except RuntimeError as exc:
        message = str(exc)
        assert "analyst-beta_report.md" in message
        assert "summary placeholder" in message


def test_run_batch_folder_reruns_failed_pdf_on_resume(tmp_path: Path, monkeypatch):
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()
    (input_dir / "case-a.pdf").write_bytes(b"%PDF-1.4")

    result_dir = output_dir / "case-a"
    result_dir.mkdir()
    result_dir.joinpath("run_status.json").write_text(
        json.dumps(
            {
                "pdf_filename": "case-a.pdf",
                "status": "failed",
                "error": "old failure",
            }
        ),
        encoding="utf-8",
    )

    calls: list[str] = []

    def fake_run_end_to_end(
        pdf_path: str,
        prompt_path: str | None = None,
        output_dir: str | None = None,
        **kwargs,
    ) -> str:
        calls.append(Path(pdf_path).name)
        result_dir = Path(output_dir)
        for report_name in (
            "analyst-alpha_report.md",
            "analyst-beta_report.md",
            "analyst-gamma_report.md",
        ):
            result_dir.joinpath(report_name).write_text(
                complete_report(7),
                encoding="utf-8",
            )
        return "ok"

    monkeypatch.setattr("dili_rucam_agents.batch.run_end_to_end", fake_run_end_to_end)

    run_batch_folder(
        input_dir=str(input_dir),
        output_dir=str(output_dir),
    )

    assert calls == ["case-a.pdf"]
    status = json.loads((result_dir / "run_status.json").read_text(encoding="utf-8"))
    assert status["status"] == "completed"
